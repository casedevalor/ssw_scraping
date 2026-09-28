import csv
import logging

from dataclasses import dataclass, field
from enum import StrEnum
from functools import cache
from itertools import islice
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple

import yaml
from pydantic import BaseModel, Field, ValidationError, model_validator

from app.utils.custom_logger import get_logger


log = get_logger()

# ---------------------------------------------------------------------------
# Baselines live in config/sentinel/<ROUTINE>.yml — one file per routine.
# A routine without a baseline file is simply not guarded.
# ---------------------------------------------------------------------------
BASELINE_DIR: Path = Path(__file__).resolve().parents[2] / "config" / "sentinel"

# When a baseline omits `header_line`, the header is auto-detected within
# this many lines at the top of the file.
HEADER_SEARCH_LINES: int = 10


# ---------------------------------------------------------------------------
# Baseline contract (validated with Pydantic)
# ---------------------------------------------------------------------------

class ColumnBaseline(BaseModel):
    """
    Expected header layout of an SSW report file.

    `columns` maps the 1-based position to the exact header name, the same
    shape used when documenting the report, so the YAML reads like the file.
    """
    routine: str
    header_line: Optional[int] = Field(default=None, ge=1)  # None -> auto-detect
    delimiter: str = ";"
    encoding: str = "windows-1252"
    columns: Dict[int, str]

    @model_validator(mode="after")
    def _check_columns(self) -> "ColumnBaseline":
        """Positions must be 1..N with no gaps; names non-empty and unique."""
        positions: List[int] = sorted(self.columns)
        if positions != list(range(1, len(positions) + 1)):
            missing: List[int] = _missing_positions(positions)
            invalid: List[int] = [pos for pos in positions if pos < 1]
            details: List[str] = []
            if missing:
                details.append(f"missing position(s) {missing}")
            if invalid:
                details.append(f"invalid position(s) {invalid}")
            raise ValueError(f"{', '.join(details)}; renumber the columns")

        names: List[str] = [name.strip() for name in self.columns.values()]
        if any(not name for name in names):
            raise ValueError("empty column name")

        duplicates: Set[str] = {name for name in names if names.count(name) > 1}
        if duplicates:
            raise ValueError(f"duplicated column name(s) {sorted(duplicates)}")

        return self

    @property
    def expected_names(self) -> List[str]:
        """Column names ordered by position."""
        return [self.columns[pos].strip() for pos in sorted(self.columns)]


def _missing_positions(positions: List[int]) -> List[int]:
    """Gaps in the 1..N numbering (e.g. [1, 2, 4] -> [3])."""
    last: int = max([len(positions), *positions]) if positions else 0
    return sorted(set(range(1, last + 1)) - set(positions))


def _load_raw_baseline(routine: str) -> dict:
    """Reads the baseline YAML without validating it."""
    with (BASELINE_DIR / f"{routine}.yml").open("r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


@cache
def load_baseline(routine: str) -> Optional[ColumnBaseline]:
    """
    Loads and validates the baseline for a routine.

    :return: The baseline, or None when the routine has no baseline file.
    :raises ValidationError: When the baseline file is malformed.
    """
    baseline_path: Path = BASELINE_DIR / f"{routine}.yml"
    if not baseline_path.is_file():
        return None

    return ColumnBaseline(**_load_raw_baseline(routine))


def has_baseline(routine: str) -> bool:
    """True when the routine is guarded by the column sentinel."""
    return (BASELINE_DIR / f"{routine}.yml").is_file()


# ---------------------------------------------------------------------------
# Comparison result
# ---------------------------------------------------------------------------

class SentinelStatus(StrEnum):
    OK    = "OK"
    DRIFT = "DRIFT"
    ERROR = "ERROR"


@dataclass(frozen=True)
class RenamedColumn:
    position: int     # 1-based position in the downloaded file
    expected: str
    actual: str


@dataclass(frozen=True)
class MovedColumn:
    name: str
    expected_position: int
    actual_position: int


@dataclass(frozen=True)
class PositionedColumn:
    name: str
    position: int


@dataclass
class ColumnReport:
    """Every difference found between the baseline and a downloaded header."""
    expected_count: int
    actual_count: int
    added: List[PositionedColumn] = field(default_factory=list)
    removed: List[PositionedColumn] = field(default_factory=list)
    renamed: List[RenamedColumn] = field(default_factory=list)
    moved: List[MovedColumn] = field(default_factory=list)
    duplicated: List[str] = field(default_factory=list)

    @property
    def status(self) -> SentinelStatus:
        has_drift: bool = any((self.added, self.removed, self.renamed, self.moved, self.duplicated))
        if has_drift or self.expected_count != self.actual_count:
            return SentinelStatus.DRIFT
        return SentinelStatus.OK


# ---------------------------------------------------------------------------
# Header reading
# ---------------------------------------------------------------------------

def _clean_row(row: List[str]) -> List[str]:
    """
    Strips every cell and drops trailing empty cells: SSW ends every line with
    the delimiter, which produces an empty last cell (pandas' 'Unnamed: N').
    An empty cell in the middle is kept, so it shows up as a difference.
    """
    cells: List[str] = [cell.strip() for cell in row]
    while cells and not cells[-1]:
        cells.pop()
    return cells


def _read_top_rows(file_path: Path, baseline: ColumnBaseline, max_rows: int) -> List[List[str]]:
    """
    Reads the first `max_rows` rows as lists of strings: csv via the baseline's
    delimiter/encoding, `.xlsx` (routines whose processing exports Excel, e.g.
    119/324) from the first sheet — delimiter/encoding don't apply there.
    """
    if file_path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook  # only needed for Excel outputs

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            return [
                ["" if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(max_row=max_rows, values_only=True)
            ]
        finally:
            workbook.close()

    with file_path.open("r", encoding=baseline.encoding, newline="") as f:
        reader = csv.reader(f, delimiter=baseline.delimiter)
        return list(islice(reader, max_rows))


def read_header(file_path: Path, baseline: ColumnBaseline) -> List[str]:
    """
    Reads only the top of the file — the report body is never loaded.

    With `header_line` set, that line is the header. Without it, the header
    is the line (within HEADER_SEARCH_LINES) sharing the most names with the
    baseline — so a title line above the header doesn't need configuring.

    :raises ValueError: When the header line is missing or can't be found.
    """
    search_lines: int = baseline.header_line or HEADER_SEARCH_LINES
    rows: List[List[str]] = [
        _clean_row(row) for row in _read_top_rows(file_path, baseline, search_lines)
    ]

    if baseline.header_line is not None:
        if len(rows) < baseline.header_line:
            raise ValueError(f"File has no line {baseline.header_line} (header): {file_path.name}")
        return rows[baseline.header_line - 1]

    expected: Set[str] = set(baseline.expected_names)
    scores: List[int] = [sum(cell in expected for cell in row) for row in rows]
    if not scores or max(scores) == 0:
        raise ValueError(f"No header found in the first {HEADER_SEARCH_LINES} lines: {file_path.name}")

    return rows[scores.index(max(scores))]


# ---------------------------------------------------------------------------
# Comparison
# ---------------------------------------------------------------------------

def _lcs_pairs(expected: List[str], actual: List[str]) -> List[Tuple[int, int]]:
    """
    Longest common subsequence between both header lists, as (i, j) index
    pairs. These are the columns that kept their relative order — the anchors
    everything else is classified around.
    """
    n, m = len(expected), len(actual)
    table: List[List[int]] = [[0] * (m + 1) for _ in range(n + 1)]

    for i in range(n - 1, -1, -1):
        for j in range(m - 1, -1, -1):
            if expected[i] == actual[j]:
                table[i][j] = table[i + 1][j + 1] + 1
            else:
                table[i][j] = max(table[i + 1][j], table[i][j + 1])

    pairs: List[Tuple[int, int]] = []
    i = j = 0
    while i < n and j < m:
        if expected[i] == actual[j]:
            pairs.append((i, j))
            i += 1
            j += 1
        elif table[i + 1][j] >= table[i][j + 1]:
            i += 1
        else:
            j += 1

    return pairs


def compare_columns(expected: List[str], actual: List[str]) -> ColumnReport:
    """
    Classifies every difference between the expected and actual headers.

    - moved:   name exists in both, but out of the relative order
    - renamed: an expected-only name and an actual-only name sitting in the
               same gap between two anchor columns (paired in order)
    - removed: expected-only names left over after pairing renames
    - added:   actual-only names left over after pairing renames

    A single inserted column shifts the absolute position of every column
    after it; that is reported once as 'added', not as N moves.
    """
    report = ColumnReport(expected_count=len(expected), actual_count=len(actual))

    expected_set: Set[str] = set(expected)
    actual_set: Set[str] = set(actual)
    report.duplicated = sorted({name for name in actual if actual.count(name) > 1})

    anchors: List[Tuple[int, int]] = _lcs_pairs(expected, actual)
    matched_actual: Set[int] = {j for _, j in anchors}

    # Common names outside the LCS changed their relative order
    matched_expected: Set[int] = {i for i, _ in anchors}
    for i, name in enumerate(expected):
        if i in matched_expected or name not in actual_set:
            continue
        actual_pos: int = next(j for j, a in enumerate(actual) if a == name and j not in matched_actual)
        report.moved.append(MovedColumn(name, expected_position=i + 1, actual_position=actual_pos + 1))

    # Walk the gaps between consecutive anchors (plus both list ends)
    boundaries: List[Tuple[int, int]] = [(-1, -1), *anchors, (len(expected), len(actual))]
    for (prev_i, prev_j), (next_i, next_j) in zip(boundaries, boundaries[1:]):
        gap_expected: List[int] = [i for i in range(prev_i + 1, next_i) if expected[i] not in actual_set]
        gap_actual: List[int] = [j for j in range(prev_j + 1, next_j) if actual[j] not in expected_set]

        for i, j in zip(gap_expected, gap_actual):
            report.renamed.append(RenamedColumn(position=j + 1, expected=expected[i], actual=actual[j]))

        paired: int = min(len(gap_expected), len(gap_actual))
        report.removed.extend(PositionedColumn(expected[i], i + 1) for i in gap_expected[paired:])
        report.added.extend(PositionedColumn(actual[j], j + 1) for j in gap_actual[paired:])

    return report


# ---------------------------------------------------------------------------
# Log formatting
# ---------------------------------------------------------------------------

def _excel_column(position: int) -> str:
    """
    Converts a 1-based position into its Excel column letter (1 -> A,
    27 -> AA), so whoever reads the log can jump straight to the column.
    """
    letters: str = ""
    while position > 0:
        position, remainder = divmod(position - 1, 26)
        letters = chr(ord("A") + remainder) + letters
    return letters


def _format_report(domain: str, routine: str, file_name: str, report: ColumnReport) -> str:
    """
    Builds a single log line. The domain log is written newest-first, so a
    multi-line entry would end up upside down — everything goes on one line.

    Positions are shown as Excel column letters: ADDED/RENAMED/MOVED-to
    point at the downloaded file, REMOVED/MOVED-from at the baseline.
    """
    delta: int = report.actual_count - report.expected_count
    delta_tag: str = f" ({delta:+d})" if delta else ""
    parts: List[str] = [
        f"{domain} SENTINEL",
        routine,
        file_name,
        f"Columns: {report.actual_count} found / {report.expected_count} expected{delta_tag}",
        str(report.status),
    ]

    if report.added:
        parts.append("ADDED: " + ", ".join(f"[{_excel_column(c.position)}] '{c.name}'" for c in report.added))
    if report.removed:
        parts.append("REMOVED: " + ", ".join(f"[{_excel_column(c.position)}] '{c.name}'" for c in report.removed))
    if report.renamed:
        parts.append("RENAMED: " + ", ".join(f"[{_excel_column(c.position)}] '{c.expected}' -> '{c.actual}'" for c in report.renamed))
    if report.moved:
        parts.append("MOVED: " + ", ".join(f"'{c.name}' {_excel_column(c.expected_position)} -> {_excel_column(c.actual_position)}" for c in report.moved))
    if report.duplicated:
        parts.append("DUPLICATED: " + ", ".join(f"'{name}'" for name in report.duplicated))

    return " | ".join(parts)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def _describe_missing(file_path: Path, routine: str) -> Optional[str]:
    """
    For a baseline with gaps in its numbering, names each missing position
    using the downloaded file's header — the line was deleted from the YAML,
    so the file is the only place that still knows which column it was.

    :return: e.g. "missing position(s) [R] [18] 'QTD VOLUMES'", or None when the
             baseline has no gaps (the error is something else).
    """
    try:
        raw: dict = _load_raw_baseline(routine)
        missing: List[int] = _missing_positions(list((raw.get("columns") or {}).keys()))
    except Exception:
        return None
    if not missing:
        return None

    try:
        # model_construct skips validation — the gaps would fail it again
        header: List[str] = read_header(file_path, ColumnBaseline.model_construct(**raw))
    except Exception:
        header = []

    return "missing position(s) " + ", ".join(
        f"[{_excel_column(pos)}] [{pos}]" + (f" '{header[pos - 1]}'" if pos <= len(header) else "")
        for pos in missing
    )


def _log_error(domain: str, routine: str, file_path: Path, detail: str) -> None:
    """Writes a sentinel ERROR entry, always on a single line."""
    single_line: str = " ".join(detail.split())
    log.log(
        level=logging.ERROR,
        domain=domain,
        message=f"{domain} SENTINEL | {routine} | {file_path.name} | {SentinelStatus.ERROR} | {single_line}",
    )


def check_file(file_path: Path, routine: str, domain: str) -> SentinelStatus:
    """
    Checks a downloaded file against its routine baseline and writes the
    verdict to the domain log (INFO when OK, WARNING on drift, ERROR when
    the file or baseline can't be read).

    Never raises: a column drift is a finding to report, not a reason to
    re-download — retrying would fetch the exact same layout.
    """
    try:
        baseline: Optional[ColumnBaseline] = load_baseline(routine)
        if baseline is None:
            raise FileNotFoundError(f"No baseline at {BASELINE_DIR / f'{routine}.yml'}")

        header: List[str] = read_header(file_path, baseline)
        report: ColumnReport = compare_columns(baseline.expected_names, header)

    except ValidationError as e:
        # Broken YAML: point at the file and the rule, one line, no Pydantic noise
        problems: str = "; ".join(err["msg"].removeprefix("Value error, ") for err in e.errors())
        _log_error(domain, routine, file_path, _describe_missing(file_path, routine) or problems)
        return SentinelStatus.ERROR

    except Exception as e:
        _log_error(domain, routine, file_path, str(e))
        return SentinelStatus.ERROR

    level: int = logging.INFO if report.status is SentinelStatus.OK else logging.WARNING
    log.log(level=level, domain=domain, message=_format_report(domain, routine, file_path.name, report))
    return report.status
